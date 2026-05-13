"""
Invest Solo — Excel Screener Export
"""

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))


import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.analysis.scoring_engine import score_universe
from src.data.sample_universe import get_universe_dataframe
from src.strategy.screener import screen_dividend_income, screen_global_best, screen_pea_value


def score_fill(score):
    if pd.isna(score):
        return PatternFill()
    s = max(0, min(100, score))
    if s >= 70:
        return PatternFill("solid", fgColor="C6EFCE")
    elif s >= 55:
        return PatternFill("solid", fgColor="FFEB9C")
    elif s >= 40:
        return PatternFill("solid", fgColor="FFC7CE")
    else:
        return PatternFill("solid", fgColor="FF9999")


def generate_excel(scored_df, output_path):
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    border = Border(*(Side(style="thin", color="D9D9D9"),) * 4)

    # Columns definition
    cols = [
        ("Ticker", 12),
        ("Name", 24),
        ("Sector", 18),
        ("Country", 8),
        ("Price", 10),
        ("Score", 8),
        ("Signal", 12),
        ("Valuation", 10),
        ("Health", 10),
        ("Profit", 10),
        ("P/E", 8),
        ("P/B", 8),
        ("ROE%", 8),
        ("Div%", 8),
        ("F-Score", 8),
        ("Z-Score", 9),
        ("Graham MoS%", 12),
        ("PEA", 8),
    ]
    col_keys = [
        "_ticker",
        "Name",
        "Sector",
        "Country",
        "Price",
        "Composite_Score",
        "Signal",
        "Valuation_Score",
        "Health_Score",
        "Profitability_Score",
        "PE",
        "PB",
        "_roe",
        "_div",
        "Piotroski_F",
        "Altman_Z",
        "Graham_MoS",
        "_pea",
    ]

    sheets_data = [
        ("All Companies", scored_df),
        ("PEA Value", screen_pea_value(scored_df)),
        ("Dividend Income", screen_dividend_income(scored_df)),
        ("Global Best", screen_global_best(scored_df)),
    ]

    for sheet_name, df in sheets_data:
        ws = wb.create_sheet(title=sheet_name)

        # Headers
        for i, (name, width) in enumerate(cols, 1):
            c = ws.cell(row=1, column=i, value=name)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border = border
            ws.column_dimensions[get_column_letter(i)].width = width

        # Data rows
        for row_idx, (ticker, row) in enumerate(df.iterrows(), 2):
            for col_idx, key in enumerate(col_keys, 1):
                if key == "_ticker":
                    val = ticker
                elif key == "_roe":
                    val = round(row.get("ROE", 0) * 100, 1)
                elif key == "_div":
                    val = round(row.get("DivYield", 0) * 100, 1)
                elif key == "_pea":
                    val = "✓" if row.get("PEA", False) else ""
                else:
                    val = row.get(key, "")

                if isinstance(val, float) and pd.isna(val):
                    val = ""
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = border
                cell.alignment = Alignment(horizontal="center" if col_idx > 2 else "left")

                if key == "Composite_Score" and isinstance(val, (int, float)):
                    cell.fill = score_fill(val)
                    cell.font = Font(bold=True)
                elif key == "Signal":
                    sig_colors = {
                        "Strong Buy": "006400",
                        "Buy": "228B22",
                        "Hold": "B8860B",
                        "Sell": "CC0000",
                        "Strong Sell": "8B0000",
                    }
                    if val in sig_colors:
                        cell.fill = PatternFill("solid", fgColor=sig_colors[val])
                        cell.font = Font(bold=True, color="FFFFFF")
                elif key in ("Valuation_Score", "Health_Score", "Profitability_Score") and isinstance(
                    val, (int, float)
                ):
                    cell.fill = score_fill(val)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb.save(output_path)
    print(f"✅ Excel saved to {output_path}")


if __name__ == "__main__":
    from pathlib import Path

    df = get_universe_dataframe()
    scored = score_universe(df)
    out_path = Path(__file__).resolve().parents[2] / "data" / "screener_results.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    generate_excel(scored, str(out_path))
